"""从二进制内容抽取纯文本：PDF / DOCX / 旧版 DOC / 纯文本类。"""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


class DocumentExtractError(Exception):
    """无法从该文件解析出可用文本。"""


def _safe_decode_utf8(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise DocumentExtractError("缺少依赖 pypdf，请 pip install pypdf") from e
    buf = io.BytesIO(data)
    reader = PdfReader(buf)
    parts: list[str] = []
    for page in reader.pages:
        t = page.extract_text() or ""
        parts.append(t)
    return "\n\n".join(parts).strip()


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise DocumentExtractError("缺少依赖 python-docx，请 pip install python-docx") from e
    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n\n".join(lines).strip()


def _looks_like_rtf(data: bytes) -> bool:
    head = data[:256].lstrip()
    return head.startswith(b"{\\rtf") or head.startswith(b"{\\urtf")


def _extract_rtf(data: bytes) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as e:
        raise DocumentExtractError("缺少依赖 striprtf，请 pip install striprtf") from e
    return (rtf_to_text(data.decode("utf-8", errors="ignore")) or "").strip()


def _find_soffice() -> str | None:
    for name in ("soffice", "soffice.exe"):
        p = shutil.which(name)
        if p:
            return p
    if sys.platform == "win32":
        for root in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ):
            if Path(root).is_file():
                return root
    return None


def _extract_doc_via_libreoffice(data: bytes) -> str | None:
    soffice = _find_soffice()
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as td:
        tdir = Path(td)
        src = tdir / "document.doc"
        src.write_bytes(data)
        outdir = tdir
        try:
            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "txt:Text",
                    str(src),
                    "--outdir",
                    str(outdir),
                ],
                check=True,
                capture_output=True,
                timeout=120,
                text=True,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
            return None
        out_txt = outdir / "document.txt"
        if not out_txt.is_file():
            return None
        return out_txt.read_text(encoding="utf-8", errors="ignore").strip()


def _extract_doc_via_word_com(path: Path) -> str | None:
    if sys.platform != "win32":
        return None
    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore
    except ImportError:
        return None
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(path.resolve()))
        try:
            text = str(doc.Content.Text or "")
        finally:
            doc.Close(False)
        return text.strip()
    except Exception:
        return None
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _extract_doc(data: bytes) -> str:
    if _looks_like_rtf(data):
        t = _extract_rtf(data)
        if t:
            return t

    t = _extract_doc_via_libreoffice(data)
    if t:
        return t

    if sys.platform == "win32":
        with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
            tmp.write(data)
            tmp_path = Path(tmp.name)
        try:
            t2 = _extract_doc_via_word_com(tmp_path)
            if t2:
                return t2
        finally:
            try:
                tmp_path.unlink(missing_ok=True)
            except OSError:
                pass

    raise DocumentExtractError(
        "未能从 .doc 提取文本。请安装 LibreOffice 并将 soffice 加入 PATH，"
        "或在 Windows 安装 Microsoft Word 后执行 pip install pywin32；"
        "也可将文件另存为 .docx 或 .pdf 后上传。"
    )


_SUFFIX_HANDLERS: dict[str, str] = {
    ".md": "text",
    ".txt": "text",
    ".markdown": "text",
    ".pdf": "pdf",
    ".docx": "docx",
    ".doc": "doc",
}


def extract_plain_text(filename: str, data: bytes) -> str:
    """
    按扩展名从二进制内容抽取纯文本。
    文本类扩展名按 UTF-8 容错解码（与原先行为一致）。
    """
    name = (filename or "unnamed.txt").strip()
    suf = Path(name).suffix.lower()
    kind = _SUFFIX_HANDLERS.get(suf)
    if kind == "text":
        return _safe_decode_utf8(data).strip()
    if kind == "pdf":
        return _extract_pdf(data)
    if kind == "docx":
        return _extract_docx(data)
    if kind == "doc":
        return _extract_doc(data)
    raise DocumentExtractError(f"不支持的扩展名: {suf or '(无)'}")


def allowed_suffixes() -> frozenset[str]:
    return frozenset(_SUFFIX_HANDLERS.keys())
